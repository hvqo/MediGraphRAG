"""Historical experiment recovered from docx/process/data/dataProcess_tex.ipynb.

Source preserved for reference, not the supported application pipeline.
Imports are safe: no model calls or data writes occur on import.
Use python -m backend.app.cli for maintained offline operations.
"""

def historical_pipeline():
    # source-cell:1
    import os  
    import docx
    import re
    from docx import Document

    # source-cell:2
    # 1.to_docx

    # source-cell:3
    import os
    import re
    from docx import Document
    from tqdm import tqdm  # 导入 tqdm

    def to_docx(main_path, save_dir, middle_path='content'):
    
        # 获取该文件夹下的所有文件和文件夹名称  
        entries = os.listdir(main_path)  
        # 过滤出文件夹名称  
        folders = [entry for entry in entries if os.path.isdir(os.path.join(main_path, entry))]  
        if '.git' in folders:
            folders.remove('.git')
        
        if not os.path.exists(save_dir):
            # 创建文件夹  
            os.makedirs(save_dir)
        
        for book_name in folders:
            path = os.path.normpath(os.path.join(main_path, book_name, middle_path))  # 获取路径
            tex_files = [filename for filename in os.listdir(path) if filename.endswith('.tex') and (filename.startswith('chap') or filename.startswith('part'))]

            # 使用 tqdm 包装 tex_files 以显示进度条
            for file in tqdm(tex_files, desc=f'Processing {book_name}'):
                full_path = os.path.join(path, file)
                # 读取文件内容
                with open(full_path, 'r', encoding='utf-8') as f:
                    content = f.read()

                # 去除多余的换行
                cleaned_content = re.sub(r'\n+', '\n\n', content).strip()

                # 将两个英文单词之间的换行替换为空格
                cleaned_content = re.sub(r'([A-Za-z]+)\n+([A-Za-z]+)', r'\1 \2', cleaned_content)
            
                # 处理"{Hp}"或"{HP}"后面可能跟有无意义换行符的情况
                cleaned_content = re.sub(r'(\{[Hh][Pp]\})\n+', r'\1', cleaned_content)

                # 创建一个新的 Word 文档
                doc = Document()

                for line in cleaned_content.splitlines():
                    if line.strip():
                        doc.add_paragraph(line)

                output_folder = f'{save_dir}/{book_name}'
                if not os.path.exists(output_folder):
                    # 创建文件夹  
                    os.makedirs(output_folder)

                doc.save(os.path.join(output_folder, f'{file[:-4]}.docx'))

            
    # main_path = '../Git_projects/medical-books/'  
    # middle_path = 'content'  # 去掉斜杠，放在最后拼接
    # save_dir = 'docxs'

    # to_docx(main_path=main_path, save_dir=save_dir, middle_path=middle_path)

    # source-cell:4
    # 2.read_docx

    # source-cell:5
    def read_docx(file_path):
        doc = Document(file_path)
        text = []
        for i in doc.paragraphs:
            text.append(i.text)
        return '\n\n'.join(text)

    # source-cell:6
    # 3.filter_content

    # source-cell:7
    def filter_content(content):
        """
        过滤掉列表中的内容，包括：
        1. 以 '\\begin' 开头并以 '\\end' 结尾的块内容（支持嵌套的情况）。
        2. 以 '\\protect' 开头的行。
    
        参数:
        content (list): 包含字符串的列表，每个字符串可能是你要处理的行或块。
    
        返回:
        list: 过滤后的列表。
        """
        filtered_content = []
        skip_block = 0  # 跟踪嵌套的 \\begin 和 \\end
        for item in content:
            # 如果遇到 \\begin，则增加嵌套层级
            if '\\begin{' in item:
                skip_block += 1
        
            # 如果在跳过状态中，直接跳过当前块内的内容
            if skip_block > 0:
                # 如果遇到 \\end，则减少嵌套层级
                if '\\end{' in item:
                    skip_block -= 1
                continue  # 跳过当前块内的内容
        
            # 如果当前行以 \\protect 开头，则跳过该行
            if item.startswith('\\protect'):
                continue
        
            # 如果不在跳过状态，则添加当前行
            filtered_content.append(item)
 
        return filtered_content 

    # source-cell:8
    # 4.to_tex

    # source-cell:9
    import os
    import re
    from docx import Document
    from tqdm import tqdm  # 导入 tqdm

    def to_tex(path, save_dir):
    
        # 获取文件夹名称  
        entries = os.listdir(path)  
        # 过滤出文件夹名称  
        if ".ipynb_checkpoints" in entries:
            entries.remove('.ipynb_checkpoints')
        
        if not os.path.exists(save_dir):
            # 创建文件夹  
            os.makedirs(save_dir)
        
        for book_name in entries:
            all_path = os.path.normpath(os.path.join(path, book_name))  # 获取单一文件夹内的所有docx根路径
            book_files = [filename for filename in os.listdir(all_path) if filename.endswith('.docx')]  # 将docx名字存入列表

            # 使用 tqdm 包装 tex_files 以显示进度条
            for file in tqdm(book_files, desc=f'Processing {book_name}'):
                full_path = os.path.join(all_path, file)  # 将根路径与docx名字拼接起来
                # 读取文件内容
                content = read_docx(full_path)
                content = content.split('\n\n')
                content = filter_content(content)
                text = '\n'.join(content)
            
            
                output_folder = f'{save_dir}/{book_name}'
                if not os.path.exists(output_folder):
                    # 创建文件夹  
                    os.makedirs(output_folder)

                with open(f'./{output_folder}/{file[:-5]}.tex', 'w', encoding='utf-8') as f:
                    f.write(text)

    # path = 'docxs'
    # save_dir = 'texs'
    # to_tex(path, save_dir)

    # source-cell:10
    # 5.read_tex

    # source-cell:11
    from pylatexenc.latex2text import LatexNodes2Text
    import re
    from collections import defaultdict
    import os
    from tqdm import tqdm

    # source-cell:12
    def read_tex(path):
        with open(path, 'r', encoding='utf-8') as f:
            content = f.read()
    
        content = LatexNodes2Text().latex_to_text(content)
        return content

    # source-cell:13
    # 6.add_content

    # source-cell:14
    def add_content(lst, book_name, chapter, section, subsection, subsubsection, content, subsubsubsection=None,):
        lst.append({
            'book_name': book_name,
            'chapter': chapter,
            'section': section,
            'subsection': subsection,
            'subsubsection': subsubsection,
            'subsubsubsection': subsubsubsection,
            'content': content,
        })
        return lst

    # source-cell:15
    # 7.parase_document

    # source-cell:16
    # 去除换行，根据以"PART"为章节或是以"CHAPTER"为章节进行针对处理
    def parse_document(book_name, text):
        # 去除所有换行符并分段
        text = text.replace('\n', ' ').strip()  # 替换所有换行符为空格
        segments = text.split('  ')  # 根据两个空格分段
        result = []
        book_name = book_name
        chapter = None
        section = None
        subsection = None
        subsubsection = None
        subsubsubsection = None  # 独属于PART
        content = None

        # 判断是否是 PART 类型章节
        is_part = segments[0].startswith("PART")
        is_section_part = True  # 判断列表内容中是否有"section"作为一级标题
    
        for idx, line in enumerate(segments):
            line = line.strip()
            # print("输出内容", line)
        
            # 处理PART章节逻辑
            if is_part:
                # for i in segments:
                #     # 如果当中出现以"section"为一级标题
                #     i = i.strip()
                #     if i.startswith("§") and not i.startswith("§.§") and not i.startswith("§.§.§"):
                #         is_section_part = True
                #         break # 跳出循环即可
                # print(is_section_part)
                # 如果有"section"
                if is_section_part:
                    if line.startswith("PART"): 
                        if idx == 0:
                            temp = segments[1].strip().split(' ')
                            chapter = temp[0]
                            if len(temp) > 1:  # 如果能两层
                                content = temp[1]
                                # print(content)
                        continue
                    elif line.startswith("CHAPTER"):  # PART下的CHAPTER为一级标题
                        if section != line.replace("CHAPTER:", "").strip():  # 如果出现了新的标题
                            if content:
                                result = add_content(
                                    lst=result, book_name = book_name, chapter=chapter, subsection=subsection, subsubsection=subsubsection
                                    , subsubsubsection=subsubsubsection, content=content, section=section
                                )
                                content = None  # 置空
                            subsection = subsubsection = subsubsubsection = None  # 清空子标题
                            
                        section = line.replace("CHAPTER:", "").strip()
                        content = None  # 一级标题下的内容清空
                        continue
                    elif line.startswith("§") and not line.startswith("§.§") and not line.startswith("§.§.§"):  # PART下的§为二级标题
                        if subsection != line.replace("§", "").strip():  # 如果出现了新的标题
                            if content:
                                result = add_content(
                                    lst=result, book_name= book_name, chapter=chapter, subsection=subsection, subsubsection=subsubsection
                                    , subsubsubsection=subsubsubsection, content=content, section=section
                                )
                                content = None  # 置空
                            subsubsection = subsubsection = None  # 清空子标题
                        
                        subsection = line.replace("§", "").strip()
                        continue
                    elif line.startswith("§.§") and not line.startswith("§.§.§"):  # PART下的§.§为三级标题
                        if subsubsection != line.replace("§.§", "").strip():
                        
                            if content:
                                result = add_content(
                                    lst=result, book_name= book_name, chapter=chapter, subsection=subsection, subsubsection=subsubsection
                                    , subsubsubsection=subsubsubsection, content=content, section=section
                                )
                                content = None  # 置空
                            subsubsubsection = None
                        
                        subsubsection = line.replace("§.§", "").strip()
                        continue
                    elif line.startswith("§.§.§"):
                        if subsubsubsection != line.replace("§.§.§", "").strip():
                            if content:
                                result = add_content(
                                        lst=result, book_name = book_name, chapter=chapter, subsection=subsection, subsubsection=subsubsection
                                        , subsubsubsection=subsubsubsection, content=content, section=section
                                    )
                                content = None  # 置空
                        
                        subsubsubsection = line.replace("§.§.§", "").strip()
                        continue
                    else:
                        # 如果正好对上了"PART"的逻辑
                        if line.split(' ')[0] == chapter:
                            continue
                        else:
                            if content:  # 如果已经有了content，则添加
                                content += line.strip()
                                continue
                            else:
                                content = line.strip()
                            

            # 处理CHAPTER章节逻辑
            else:
                if line.startswith("CHAPTER"):
                    chapter = line.replace("CHAPTER:", "").strip()
                    content = None  # 章节内容清空
                    continue
                elif line.startswith("§") and not line.startswith("§.§") and not line.startswith("§.§.§"):  # CHAPTER下的§为一级标题
                    if section != line.replace("§", "").strip():  # 如果出现了新的一级标题则清空二级标题和三级标题
                        if content:
                            result = add_content(
                                    lst=result, book_name= book_name, chapter=chapter, subsection=subsection, subsubsection=subsubsection
                                    , subsubsubsection=subsubsubsection, content=content, section=section
                                )
                            content = None  # 置空
                        subsection = None
                        subsubsection = None
                    
                    section = line.replace("§", "").strip()
                    content = None  # 一级标题下的内容清空
                    continue
                elif line.startswith("§.§") and not line.startswith("§.§.§"):  # CHAPTER下的§.§为二级标题
                    if subsection != line.replace("§.§", "").strip():  # 如果出现了新的二级标题则清空二级标题和三级标题
                        if content:
                            result = add_content(
                                    lst=result, book_name= book_name, chapter=chapter, subsection=subsection, subsubsection=subsubsection
                                    , subsubsubsection=subsubsubsection, content=content, section=section
                                )
                            content = None  # 置空
                        subsubsection = None
                    
                    subsection = line.replace("§.§", "").strip()
                    content = None
                    continue
                elif line.startswith("§.§.§"):  # CHAPTER下的§.§.§为三级标题
                    if subsubsection != line.replace("§.§.§", "").strip():
                        if content:
                            result = add_content(
                                    lst=result, book_name= book_name, chapter=chapter, subsection=subsection, subsubsection=subsubsection
                                    , subsubsubsection=subsubsubsection, content=content, section=section
                                )
                            content = None  # 置空
                        
                    subsubsection = line.replace("§.§.§", "").strip()
                    continue
                else:
                    if content:
                        content += line.strip()
                        continue
                    else:
                        content = line.strip()

        return result

    # source-cell:17
    # 8.know_extract

    # source-cell:18
    folder_path = './texs'
    book_path = os.listdir(folder_path)
    path = []
    for i in book_path:
        path.append(os.path.normpath(os.path.join(folder_path, i)))
    
    lst = []

    for book in path:
        tqdm_name = book.split('/')[1]
        files = [filename for filename in os.listdir(book) if filename.endswith('.tex')]
    
        for filename in tqdm(files, desc=f'Procession book: {tqdm_name}'):
            file_path = os.path.join(book, filename)
        
            content = read_tex(file_path)
            content = parse_document(tqdm_name, content)
            lst.extend(content)
        
    content = lst

    # source-cell:19
    from docx import Document

    from langchain_community.llms import Ollama
    from langchain_core.output_parsers import StrOutputParser,JsonOutputParser
    from langchain_core.prompts import ChatPromptTemplate,MessagesPlaceholder
    from langchain.memory import ConversationBufferMemory
    from langchain_core.chat_history import BaseChatMessageHistory
    import pandas as pd

    import os
    from docx import Document
    import json
    from tqdm import tqdm  
    import concurrent.futures


    import warnings
    warnings.filterwarnings("ignore")

    # source-cell:20
    # 确保输出是有效的JSON格式
    systemContent = r"""你是我的医学知识提取助理，负责从医学文件中提取结构化的医学知识点。每条知识点应围绕同一医学概念，将密切关联的知识点合并为一个独立的完整条目。

    要求如下：
    - 输入内容为一个包含以下字段的字典：'book_name'（书名）、'chapter'（章节）、'section'（一级标题）、'subsection'（二级标题）、'subsubsection'（三级标题）、'subsubsubsection'（四级标题）、'content'（内容）。
    - 提取具体的医学知识点，专注于病理学定义、影像学特征、临床症状等，不包含推断、诊治内容或流行病学数据。
    - 若存在多个关联性强的知识点（例如围绕 GERD、GU/DU 等主题），应将它们合并为一个结构化条目，保持内容的连贯性。
    - 每条提取的知识点应独立且完整，避免模糊表达，但内容相近的知识点应整合在一起形成逻辑完整的信息。
    - 输出格式为 JSON 列表，其中每个条目为一个字典，包含一条知识点信息。
    - 输出中应仅包含 'knowledge' 字段，忽略输入中的其他标题字段。

    输入案例:
    {knowledge}

    输出案例：
    [
        {{"knowledge": "中耳胆脂瘤是指在中耳腔内形成的囊性病变，通常由表皮细胞异常增生引起。MRI影像学特征包括：T1加权像（T1WI）呈中等信号强度，T2加权像（T2WI）显示不均匀高信号，并且弥散加权成像（DWI）上病灶的弥散受限明显，增强扫描无强化表现。"}},
        {{"knowledge": "GERD相关症状包括声音嘶哑、夜间睡眠障碍、咽炎、耳痛和龈炎等。这些症状可能与胃酸反流及其对周围组织的影响有关。"}}
    ]
    """

    prompt_template = ChatPromptTemplate.from_messages(
        [("system", systemContent), ("user", "{text}")]
    )

    model = Ollama(model="qwen2.5:14b-32k", temperature=0.0)
    parser = JsonOutputParser()
    chain = prompt_template | model | parser

    # source-cell:21
    # 确保文件夹存在
    os.makedirs('./know', exist_ok=True)

    if not os.path.exists('./know/knowledge.json'):
        with open('./know/knowledge.json', 'w') as f:
            json.dump([], f)

    # 读取现有数据
    with open('./know/knowledge.json', 'r') as f:
        responses = json.load(f)

    for i in tqdm(content):
        time = 0
        generate_time = 0
        while True:
            try:     
                knowledge = {
                    'book_name': i['book_name'],
                    'chapter': i['chapter'],
                    'section': i['section'],
                    'subsection': i['subsection'],
                    'subsubsection': i['subsubsection'],
                    'subsubsubsection': i['subsubsubsection'],
                    'content': i['content']
                }
                response = chain.invoke({'knowledge': knowledge, 'text': knowledge})
                if isinstance(response, list):
                    if response:
                        # print(response)
                        for i in response:
                            responses.append(i)
                    
                        # 立即写入文件
                        with open('./know/knowledge.json', 'w') as f:
                            json.dump(responses, f, ensure_ascii=False, indent=4)
                        break
                    else:
                        break
                else:
                    generate_time += 1
                    if generate_time == 50:
                        break
            except Exception as e:
                time+=1
                print(str(e))
                if time == 5:
                    b = i["book_name"]
                    c = i["chapter"]
                    sec = i["section"]
                    subsec = i["subsection"]
                    subsubsec = i["subsubsection"]
                    subsubsubsec = i["subsubsubsection"]
                    print(f"{b}_{c}_{sec}_{subsec}_{subsubsec}_{subsubsubsec} happen error")
                    break

    # source-cell:22



if __name__ == "__main__":
    print("Historical source archive. Use python -m backend.app.cli --help for supported commands.")
