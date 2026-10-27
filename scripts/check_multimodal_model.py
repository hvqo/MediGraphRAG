"""Real local vision + generation smoke test on a public abstract, with isolated personal data."""

import io
import json
from pathlib import Path
import re
import sys
import tempfile
import textwrap
import time
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from backend.app.main import create_app


def edit_distance(a, b):
    row = list(range(len(b) + 1))
    for i, x in enumerate(a, 1):
        nxt = [i]
        for j, y in enumerate(b, 1):
            nxt.append(min(nxt[-1] + 1, row[j] + 1, row[j - 1] + (x != y)))
        row = nxt
    return row[-1]


def main():
    raw = json.loads((ROOT / ".runtime/medical-data/ori_pqal.json").read_text(encoding="utf-8"))
    pmid, record = next(
        (i, r)
        for i, r in raw.items()
        if re.search(r"diabet|hypertens|cardiac", r["QUESTION"], re.I) and 150 < len(r["CONTEXTS"][0]) < 700
    )
    text = record["CONTEXTS"][0]
    image = Image.new("RGB", (1100, 650), "white")
    draw = ImageDraw.Draw(image)
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    if not font_path.exists():
        font_path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    font = ImageFont.truetype(str(font_path), 26)
    draw.multiline_text((35, 35), "\n".join(textwrap.wrap(text, 72)), fill="black", font=font, spacing=12)
    png = io.BytesIO()
    image.save(png, format="PNG")
    pdf = io.BytesIO()
    image.save(pdf, format="PDF", resolution=120)
    word = Document()
    word.add_paragraph(text)
    word.add_picture(io.BytesIO(png.getvalue()), width=Inches(6))
    docx = io.BytesIO()
    word.save(docx)
    report = {
        "source": "https://github.com/pubmedqa/pubmedqa",
        "pmid": pmid,
        "ocr": "RapidOCR 1.4.4 CPU text-first, vision fallback",
        "vision_model": "qwen2.5vl:7b",
        "generation_model": "qwen2.5",
        "scope": "One clean machine-printed abstract rendered to PNG/scanned PDF/embedded Word image; functional smoke test, not representative OCR or clinical accuracy. All original text stays in ignored runtime.",
        "files": [],
    }
    with tempfile.TemporaryDirectory(prefix="multimodal-") as folder, TestClient(create_app(folder)) as client:
        client.put("/settings", json={"request_timeout": 600})
        start = time.perf_counter()
        response = client.post(
            "/workspace/submit",
            data={"instruction": "根据这些医学资料生成1道练习题，考察原文明确提到的研究内容。"},
            files=[
                ("files", ("medical.png", png.getvalue())),
                ("files", ("scanned.pdf", pdf.getvalue())),
                ("files", ("medical.docx", docx.getvalue())),
            ],
        )
        response.raise_for_status()
        identifier = response.json()["id"]
        while True:
            row = client.get("/workspace/operations/" + identifier).json()
            print(row["status"], row["progress"]["stage"], flush=True)
            if row["status"] in ["foreground_done", "completed", "failed"]:
                break
            if time.perf_counter() - start > 1200:
                raise TimeoutError(row["progress"])
            time.sleep(5)
        report["foreground_seconds"] = time.perf_counter() - start
        report["actual_vision_calls"] = sum(t["role"] == "vision" for t in client.get("/traces").json())
        report["request_status"] = row["result"].get("request_status")
        report["error"] = row["error"] or row["result"].get("question_error")
        report["generated_questions"] = len(row["result"].get("question_ids", []))
        normalize = lambda s: re.sub(r"[^a-z0-9]", "", s.lower())
        for item in row["result"].get("files", []):
            entry = {k: v for k, v in item.items() if k in ["filename", "status", "error"]}
            if item.get("id"):
                doc = client.get("/documents/" + item["id"]).json()
                candidate = "\n".join(c["text"] for c in doc["chunks"])
                if item["filename"].endswith(".docx"):
                    candidate = "\n".join(c["text"] for c in doc["chunks"] if c["location"].startswith("内嵌图"))
                candidate = candidate.split("[视觉模型描述")[0].strip()
                entry["character_error_rate"] = edit_distance(normalize(text), normalize(candidate)) / len(
                    normalize(text)
                )
                entry["chunks"] = len(doc["chunks"])
            report["files"].append(entry)
        if row["result"].get("document_ids"):
            t = time.perf_counter()
            answer = client.post(
                "/ask", json={"query": record["QUESTION"], "document_ids": row["result"]["document_ids"]}
            )
            report["single_question_api_seconds"] = time.perf_counter() - t
            report["single_question_http_status"] = answer.status_code
            report["single_question_answer_status"] = answer.json().get("status")
        deadline = time.monotonic() + 180
        while row["status"] == "foreground_done" and time.monotonic() < deadline:
            time.sleep(2)
            row = client.get("/workspace/operations/" + identifier).json()
        report["archive_status"] = row["status"]
        report["background"] = row["result"].get("background")
        report["index_ready"] = client.get("/indexes/status").json()["ready"]
        report["graph_relations"] = len(client.get("/relations").json())
        output = ROOT / "benchmarks/results/multimodal-real.json"
        output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(report, ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    main()
